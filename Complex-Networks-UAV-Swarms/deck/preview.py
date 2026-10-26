"""Kaba önizleme: .pptx → PNG (PIL + Arial ölçüleri). LibreOffice/Keynote gerektirmez.

Gerçek render değildir; amacı metin taşmasını ve yerleşimi görmek. Taşan metin kutularını
stderr'e yazar. Kullanım: python3 deck/preview.py deck/out/Sunum1_IHA_Suruleri.pptx
"""

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

DPI = 100
FONTS = Path("/System/Library/Fonts/Supplemental")
_cache: dict[tuple[bool, int], ImageFont.FreeTypeFont] = {}


def px(emu) -> int:
    return round(emu / 914400 * DPI)


def font(size_pt: float, bold: bool) -> ImageFont.FreeTypeFont:
    key = (bold, round(size_pt * DPI / 72))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(str(FONTS / ("Arial Bold.ttf" if bold else "Arial.ttf")), key[1])
    return _cache[key]


def rgb(color, default=(29, 36, 48)):
    try:
        return tuple(color.rgb)
    except Exception:
        return default


def layout_paragraph(p, width):
    """Paragrafı kelime kelime sarar → [(satır parçaları, satır yüksekliği)], madde imi girintisi."""
    ppr = p._p.pPr
    indent = px(int(ppr.get("marL", 0))) if ppr is not None else 0
    bullet = ppr is not None and ppr.find(qn("a:buChar")) is not None
    words = []
    for r in p.runs:
        size = r.font.size.pt if r.font.size else 18
        sub = r.font._element.get("baseline") is not None
        f = font(size * (0.7 if sub else 1), bool(r.font.bold))
        color = rgb(r.font.color)
        for k, w in enumerate(r.text.split(" ")):
            words.append((w, f, color, sub, size, k > 0))
    lines, cur, cur_w = [], [], 0
    avail = width - indent
    for w, f, color, sub, size, spaced in words:
        ww = f.getlength(w)
        sp = f.getlength(" ") if (spaced and cur) else 0
        if cur and cur_w + sp + ww > avail:
            lines.append(cur)
            cur, cur_w, sp = [], 0, 0
        cur.append((w, f, color, sub, size, sp))
        cur_w += sp + ww
    if cur:
        lines.append(cur)
    out = []
    for ln in lines:
        size = max(s for *_, s, _ in ln)
        out.append((ln, size * DPI / 72 * 1.2 * (p.line_spacing or 1)))
    return out, indent, bullet


def draw_text(draw, tf, x, y, w, h, name, problems):
    ml, mr = px(tf.margin_left), px(tf.margin_right)
    mt, mb = px(tf.margin_top), px(tf.margin_bottom)
    inner_w = w - ml - mr
    blocks, total = [], 0
    for i, p in enumerate(tf.paragraphs):
        lines, indent, bullet = layout_paragraph(p, inner_w)
        before = (p.space_before.pt * DPI / 72) if (i and p.space_before) else 0
        blocks.append((p, lines, indent, bullet, before))
        total += before + sum(lh for _, lh in lines)
    if total > h - mt - mb + 2 and any(p.runs for p in tf.paragraphs):
        problems.append(f"{name}: metin {total:.0f}px, kutu {h - mt - mb}px")
    anchor = tf.vertical_anchor
    cy = y + mt
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y + (h - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - total
    for p, lines, indent, bullet, before in blocks:
        cy += before
        for li, (ln, lh) in enumerate(lines):
            line_w = sum(f.getlength(t) + sp for t, f, _, _, _, sp in ln)
            cx = x + ml + indent
            if p.alignment == PP_ALIGN.CENTER:
                cx = x + ml + (inner_w - line_w) / 2
            elif p.alignment == PP_ALIGN.RIGHT:
                cx = x + w - mr - line_w
            if bullet and li == 0:
                f0 = ln[0][1]
                draw.text((x + ml, cy), "■", font=font(ln[0][4] * 0.6, False), fill=(47, 111, 235))
            for t, f, color, sub, size, sp in ln:
                cx += sp
                draw.text((cx, cy + (size * DPI / 72 * 0.45 if sub else 0)), t, font=f, fill=color)
                cx += f.getlength(t)
            cy += lh


def render(prs_path: Path, out_dir: Path) -> list[str]:
    prs = Presentation(prs_path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    out_dir.mkdir(parents=True, exist_ok=True)
    problems, images = [], []
    for n, slide in enumerate(prs.slides, start=1):
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img.paste(Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((w, h)), (x, y))
                continue
            if sh.has_table:
                tbl = sh.table
                ry = y
                for row in tbl.rows:
                    rx = x
                    for j, cell in enumerate(row.cells):
                        cw = px(tbl.columns[j].width)
                        fill = rgb(cell.fill.fore_color, (255, 255, 255)) if cell.fill.type == MSO_FILL.SOLID else None
                        d.rectangle([rx, ry, rx + cw, ry + px(row.height)], fill=fill, outline=(216, 222, 233))
                        draw_text(d, cell.text_frame, rx, ry, cw, px(row.height), f"slayt {n} tablo", problems)
                        rx += cw
                    ry += px(row.height)
                continue
            fill = None
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                if sh.fill.type == MSO_FILL.SOLID:
                    fill = rgb(sh.fill.fore_color)
                elif sh.fill.type == MSO_FILL.GRADIENT:
                    fill = rgb(sh.fill.gradient_stops[0].color)
                outline, lw = None, 0
                if sh.line.fill.type == MSO_FILL.SOLID:
                    outline, lw = rgb(sh.line.color), max(1, px(sh.line.width))
                kind = sh.auto_shape_type
                if kind == MSO_SHAPE.OVAL:
                    d.ellipse([x, y, x + w, y + h], fill=fill, outline=outline, width=lw)
                elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) // 4, fill=fill, outline=outline,
                                        width=lw)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=fill, outline=outline, width=lw)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text(d, sh.text_frame, x, y, w, h, f"slayt {n} '{sh.text_frame.text[:25]}'", problems)
        path = out_dir / f"slide-{n:02d}.png"
        img.save(path)
        images.append(img)
    # 3 sütunlu temas sayfası
    cols, scale = 3, 0.33
    tw, th = int(W * scale), int(H * scale)
    rows = (len(images) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw + (cols + 1) * 10, rows * th + (rows + 1) * 10), (200, 200, 200))
    for i, im in enumerate(images):
        sheet.paste(im.resize((tw, th)), (10 + (i % cols) * (tw + 10), 10 + (i // cols) * (th + 10)))
    sheet.save(out_dir / "contact.png")
    return problems


if __name__ == "__main__":
    src = Path(sys.argv[1])
    for msg in render(src, src.parent / "preview"):
        print("TAŞMA:", msg, file=sys.stderr)
    print(src.parent / "preview" / "contact.png")
