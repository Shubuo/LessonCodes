"""Sunum 1 destesi: deck/sunum1.yaml → deck/out/Sunum1_IHA_Suruleri.pptx.

Projeksiyon kuralları (PLAN.md): Arial; başlık 32 pt, gövde ≥ 22 pt, tablo ≥ 18 pt,
kaynak notu 12 pt; açık zemin + koyu metin; kenar boşluğu ≥ 0,5 inç.
Metin biçimi: _{...} alt simge, ^{...} üst simge, **...** kalın.
"""

import re
import shutil
import sys
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
CONTENT = Path(__file__).parent / "sunum1.yaml"
OUT = Path(__file__).parent / "out"
DRIVE = Path("/Users/buraky/Library/CloudStorage/GoogleDrive-burakyoruk94@gmail.com/My Drive/"
             "2_ACADEMIA/PhD /2026-GUZ/Alg4ComplexNtw/Sunum1")

FONT = "Arial"
NAVY = RGBColor(0x17, 0x32, 0x4D)
BLUE = RGBColor(0x2F, 0x6F, 0xEB)
CYAN = RGBColor(0x5E, 0xC2, 0xD1)
TEXT = RGBColor(0x1D, 0x24, 0x30)
MUTED = RGBColor(0x66, 0x70, 0x85)
SOFT = RGBColor(0xF5, 0xF7, 0xFB)
BORDER = RGBColor(0xD8, 0xDE, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ALERT = RGBColor(0xD5, 0x5E, 0x00)

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)
CW = W - 2 * M                      # içerik genişliği
TOP = Inches(1.45)                  # başlığın altı
BOTTOM = Inches(6.5)                # kaynak notunun üstü
NUMBERED = 16                       # kaynakça numarasız

MARKUP = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


# ---------- metin yardımcıları ----------

def add_runs(p, text: str, size: float, color=TEXT, bold=False):
    """`text` içindeki **kalın**, _{alt} ve ^{üst} işaretlerini run'lara çevirir."""
    for part in MARKUP.split(text):
        if not part:
            continue
        run = p.add_run()
        run_bold, baseline = bold, None
        if part.startswith("**"):
            part, run_bold = part[2:-2], True
        elif part.startswith("_{"):
            part, baseline = part[2:-1], "-25000"
        elif part.startswith("^{"):
            part, baseline = part[2:-1], "30000"
        run.text = part
        f = run.font
        f.name, f.size, f.bold = FONT, Pt(size), run_bold
        f.color.rgb = color
        if baseline:
            f._element.set("baseline", baseline)


def text_box(slide, left, top, width, height, lines, size, *, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(spacing)
        add_runs(p, line, size, color, bold)
    return box


def bullets(slide, left, top, width, height, items, size=24, gap=14):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(gap if i else 0)
        p.line_spacing = 1.05
        indent = Emu(Pt(size) * 0.95)
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(indent))
        ppr.set("indent", str(-indent))
        clr = ppr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": "2F6FEB"}))
        ppr.append(clr)
        ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": FONT}))
        ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "■"}))
        add_runs(p, item, size)
    return box


def rect(slide, left, top, width, height, fill, *, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def picture_fit(slide, path: Path, left, top, width, height):
    """Oranı koruyarak kutuya sığdırır ve ortalar."""
    w_px, h_px = Image.open(path).size
    scale = min(width / w_px, height / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    return slide.shapes.add_picture(str(path), left + (width - w) // 2, top + (height - h) // 2, w, h)


# ---------- ortak çerçeve ----------

def accent_bar(slide):
    bar = rect(slide, 0, 0, W, Inches(0.09), NAVY)
    bar.fill.gradient()
    bar.fill.gradient_angle = 0
    stops = bar.fill.gradient_stops
    stops[0].color.rgb, stops[0].position = NAVY, 0
    stops[1].color.rgb, stops[1].position = CYAN, 1.0


def footer(slide, sections, section, number):
    box = slide.shapes.add_textbox(M, Inches(7.0), Inches(9.5), Inches(0.3))
    tf = box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for i, name in enumerate(sections):
        run = p.add_run()
        run.text = name + ("   ·   " if i < len(sections) - 1 else "")
        f = run.font
        f.name, f.size, f.bold = FONT, Pt(14), name == section
        f.color.rgb = BLUE if name == section else MUTED
    if number:
        text_box(slide, W - M - Inches(1.5), Inches(7.0), Inches(1.5), Inches(0.3), f"{number} / {NUMBERED}", 14,
                 color=MUTED, align=PP_ALIGN.RIGHT)


def base_slide(prs, meta, spec, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_bar(slide)
    text_box(slide, M, Inches(0.38), CW, Inches(0.9), spec["title"], 32, color=NAVY, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    footer(slide, meta["sections"], spec.get("section"), number)
    if spec.get("source"):
        text_box(slide, M, Inches(6.62), CW, Inches(0.3), spec["source"], 12, color=MUTED)
    slide.notes_slide.notes_text_frame.text = spec.get("notes", "").strip()
    return slide


def takeaway(slide, pair, top=Inches(5.88)):
    box = rect(slide, M, top, CW, Inches(0.6), SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box.adjustments[0] = 0.15
    tf = box.text_frame
    tf.margin_left = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    add_runs(p, pair[0], 22, ALERT, bold=True)
    add_runs(p, pair[1], 22, TEXT)


# ---------- düzenler ----------

def cover(prs, meta, spec, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, Inches(4.6), NAVY)
    rect(slide, 0, Inches(4.6), W, Inches(0.09), CYAN)
    text_box(slide, M, Inches(1.2), Inches(11.5), Inches(1.9), meta["title"], 44, color=WHITE, bold=True)
    text_box(slide, M, Inches(3.3), Inches(11.5), Inches(0.6), meta["subtitle"], 26, color=CYAN)
    text_box(slide, M, Inches(5.05), Inches(11.5), Inches(0.5), meta["author"], 26, bold=True)
    text_box(slide, M, Inches(5.65), Inches(11.5), Inches(1.0), [meta["course"], meta["place_date"]], 20,
             color=MUTED, spacing=4)
    slide.notes_slide.notes_text_frame.text = spec.get("notes", "").strip()


def cards(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    items = spec["cards"]
    gap = Inches(0.25)
    cw = (CW - gap * (len(items) - 1)) // len(items)
    top, ch = Inches(1.75), Inches(3.6)
    for i, card in enumerate(items):
        x = M + i * (cw + gap)
        rect(slide, x, top, cw, ch, SOFT)
        rect(slide, x, top, cw, Inches(0.1), ALERT if card.get("highlight") else NAVY)
        text_box(slide, x + Inches(0.25), top + Inches(0.4), cw - Inches(0.5), Inches(1.0), card["title"], 24,
                 color=ALERT if card.get("highlight") else NAVY, bold=True)
        text_box(slide, x + Inches(0.25), top + Inches(1.5), cw - Inches(0.5), Inches(1.9), card["text"], 22)
    takeaway(slide, spec["takeaway"], top=Inches(5.75))


def figure(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    bottom = Inches(5.75) if spec.get("takeaway") else BOTTOM
    picture_fit(slide, ROOT / spec["figure"], M, TOP, CW, bottom - TOP)
    if spec.get("takeaway"):
        takeaway(slide, spec["takeaway"])


def figure_bullets(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    picture_fit(slide, ROOT / spec["figure"], M, TOP, CW, Inches(3.75))
    half = (len(spec["bullets"]) + 1) // 2
    col_w = (CW - Inches(0.4)) // 2
    for k, chunk in enumerate((spec["bullets"][:half], spec["bullets"][half:])):
        bullets(slide, M + k * (col_w + Inches(0.4)), Inches(5.3), col_w, Inches(1.2), chunk, size=22, gap=8)


def timeline(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    items = spec["items"]
    col = CW // len(items)
    line_y = Inches(2.75)
    rect(slide, M + col // 2, line_y - Inches(0.03), col * (len(items) - 1), Inches(0.06), BORDER)
    for i, it in enumerate(items):
        x = M + i * col
        cx = x + col // 2
        text_box(slide, x, Inches(1.75), col, Inches(0.5), it["year"], 28, color=BLUE, bold=True,
                 align=PP_ALIGN.CENTER)
        dot = rect(slide, cx - Inches(0.16), line_y - Inches(0.16), Inches(0.32), Inches(0.32), NAVY,
                   shape=MSO_SHAPE.OVAL)
        dot.line.color.rgb, dot.line.width = WHITE, Pt(3)
        text_box(slide, x + Inches(0.1), Inches(3.2), col - Inches(0.2), Inches(0.5), f"[{i + 1}]", 20,
                 color=MUTED, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, x + Inches(0.1), Inches(3.65), col - Inches(0.2), Inches(1.0), it["theme"], 24,
                 color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, x + Inches(0.1), Inches(4.65), col - Inches(0.2), Inches(0.9), it["venue"], 18,
                 color=TEXT, align=PP_ALIGN.CENTER)
    takeaway(slide, spec["takeaway"])


def paper(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    chip = rect(slide, M, Inches(1.4), Inches(5.9), Inches(0.5), SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.adjustments[0] = 0.5
    tf = chip.text_frame
    tf.margin_left = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_runs(tf.paragraphs[0], spec["chip"], 18, BLUE, bold=True)
    bullets(slide, M, Inches(2.2), Inches(5.9), Inches(4.2), spec["bullets"], size=24, gap=16)
    picture_fit(slide, ROOT / spec["figure"], Inches(6.85), TOP, W - M - Inches(6.85), BOTTOM - TOP)


def table(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    cols, rows = spec["columns"], spec["rows"]
    symbols = spec.get("symbols", False)
    row_h = Inches(0.62)
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), M, Inches(1.6), CW, row_h * (len(rows) + 1))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, w in enumerate(spec["widths"]):
        tbl.columns[j].width = Inches(w)
    for i in range(len(rows) + 1):
        tbl.rows[i].height = row_h
        for j in range(len(cols)):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.15)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            header = i == 0
            cell.fill.fore_color.rgb = NAVY if header else (SOFT if i % 2 == 0 else WHITE)
            text = cols[j] if header else rows[i - 1][j]
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (symbols and j > 0) else PP_ALIGN.LEFT
            strong = (not header) and (i - 1) in spec.get("highlight_rows", [])
            if header:
                add_runs(p, text, 20, WHITE, bold=True)
            elif symbols and j > 0:
                add_runs(p, text, 26, NAVY if text == "●" else MUTED, bold=False)
            else:
                add_runs(p, text, 20, BLUE if strong else TEXT, bold=strong or j == 0)
    if spec.get("takeaway"):
        takeaway(slide, spec["takeaway"])


def gap(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, number)
    q = spec["quadrants"]
    lab_w, qw, qh = Inches(1.55), Inches(2.25), Inches(1.75)
    x0, y0 = M + lab_w, Inches(1.65)
    cells = {(0, 0): q["top_left"], (0, 1): q["top_right"], (1, 0): q["bottom_left"], (1, 1): q["bottom_right"]}
    for (r, c), label in cells.items():
        x, y = x0 + c * (qw + Inches(0.1)), y0 + r * (qh + Inches(0.1))
        target = (r, c) == (0, 1)
        box = rect(slide, x, y, qw, qh, WHITE if target else SOFT, line=ALERT if target else None)
        if target:
            box.line.width = Pt(3)
            box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, label or "—", 26, ALERT if target else (NAVY if label else MUTED), bold=True)
    for r, label in enumerate(reversed(spec["axis_y"])):
        text_box(slide, M, y0 + r * (qh + Inches(0.1)), lab_w - Inches(0.15), qh, label, 18, color=MUTED,
                 bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    for c, label in enumerate(spec["axis_x"]):
        text_box(slide, x0 + c * (qw + Inches(0.1)), y0 + 2 * qh + Inches(0.25), qw, Inches(0.7), label, 18,
                 color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    right = x0 + 2 * qw + Inches(0.55)
    rw = W - M - right
    qbox = rect(slide, right, Inches(1.65), rw, Inches(2.1), SOFT)
    rect(slide, right, Inches(1.65), Inches(0.1), Inches(2.1), ALERT)
    text_box(slide, right + Inches(0.3), Inches(1.8), rw - Inches(0.45), Inches(1.9), spec["question"], 22,
             color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    bullets(slide, right, Inches(4.05), rw, Inches(2.3), spec["bullets"], size=22, gap=12)


def references(prs, meta, spec, number):
    slide = base_slide(prs, meta, spec, None)
    text_box(slide, M, TOP, CW, BOTTOM - TOP, spec["items"], 16, spacing=10)


LAYOUTS = {"cover": cover, "cards": cards, "figure": figure, "figure_bullets": figure_bullets,
           "timeline": timeline, "paper": paper, "table": table, "gap": gap, "references": references}


# ---------- kontrol ----------

def check(prs) -> list[str]:
    """Render olmadan geometri ve punto kontrolü: taşma, kenar boşluğu, küçük punto."""
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.width == W or sh.top == 0:          # tam genişlikli şerit ve bantlar
                continue
            if sh.left < Inches(0.5) or sh.left + sh.width > W - Inches(0.5) or sh.top + sh.height > H - Inches(0.15):
                problems.append(f"slayt {n}: '{sh.name}' kenar boşluğu dışında")
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    size = r.font.size.pt if r.font.size else 0
                    small_ok = sh.top >= Inches(6.6) or n == len(prs.slides)   # kaynak notu, alt şerit, kaynakça
                    if size and size < 18 and not small_ok:
                        problems.append(f"slayt {n}: {size:.0f} pt metin '{r.text[:30]}'")
    return problems


def build(path: Path) -> Path:
    data = yaml.safe_load(CONTENT.read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for i, spec in enumerate(data["slides"], start=1):
        LAYOUTS[spec["layout"]](prs, data["meta"], spec, i)
    problems = check(prs)
    for msg in problems:
        print("UYARI:", msg, file=sys.stderr)
    OUT.mkdir(exist_ok=True)
    prs.save(path)
    return path


if __name__ == "__main__":
    out = build(OUT / "Sunum1_IHA_Suruleri.pptx")
    print(out)
    if "--drive" in sys.argv:
        DRIVE.mkdir(exist_ok=True)
        print(shutil.copy2(out, DRIVE / out.name))
